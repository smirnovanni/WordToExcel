import pandas as pd
import os
from docx.api import Document
from source.config import INPUT_DIRECTORY, OUTPUT_DIRECTORY
from openpyxl import load_workbook


def get_table(table):
    data = []
    keys = None
    for i, row in enumerate(table.rows):
        if i == 0:
            keys = tuple((r + 1 for r, cell in enumerate(row.cells)))
            continue

        text = (cell.text for cell in row.cells)

        row_data = dict(zip(keys, text))
        data.append(row_data)

    return pd.DataFrame(data)


def write_table(df, file, sheet):
    writer = pd.ExcelWriter(file, engine='xlsxwriter')
    df.to_excel(writer, sheet_name='{}_{}'.format('Sheet', sheet))
    writer.save()


def main():
    for file in os.listdir(INPUT_DIRECTORY):
        if file.endswith(".doc") or file.endswith(".docx"):
            document = Document(os.path.join(INPUT_DIRECTORY, file))
            writer = pd.ExcelWriter('{}/{}.{}'.format(OUTPUT_DIRECTORY, os.path.splitext(file)[0], 'xlsx'),
                                    engine='openpyxl')
            for i, table in enumerate(document.tables):
                df = get_table(table)
                df.to_excel(writer, sheet_name='{}{}'.format('Sheet_', i + 1), index=False)
                writer.save()

            writer.close()


if __name__ == "__main__":
    main()
